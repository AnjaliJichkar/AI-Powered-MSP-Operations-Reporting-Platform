"""The render gate: a deck with any placeholder left in it is not written.

The template marks its instructions and sample content in red (`FF0000`), and
that red can be set in three places, all of which PowerPoint honours or a
future edit of the template might rely on (slide map, rule 2):

1. on the run itself (`a:r/a:rPr/a:solidFill`);
2. as the paragraph's default (`a:p/a:pPr/a:defRPr/a:solidFill`);
3. in the shape's list style for the paragraph's level
   (`a:txBody/a:lstStyle/a:lvlNpPr/a:defRPr/a:solidFill`).

`effective_colour` resolves them in that order, for text in shapes, groups and
table cells alike, by walking the slide XML rather than the shape tree.

Some sample content is not red: the dashboard screenshot on slide 6, the CX
scores on slide 9, the sample takeaways in slide 7's table. Those are checked
by what they are, not by colour.

The gate reads the presentation that is about to be saved. It trusts nothing
the renderer says about what it filled in.
"""

from __future__ import annotations

import re
from collections.abc import Callable, Mapping
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

TEMPLATE = Path(__file__).resolve().parents[3] / "docs" / "reference" / "tbr-template.pptx"

RED = "FF0000"

#: From the slide map, rule 6. `XX` is matched as written; the rest ignore case.
FORBIDDEN_STRINGS: tuple[str, ...] = (
    "XX",
    "ABC Company",
    "202x",
    "Insert",
    "{",
    "<<",
    "talk track",
    "Examples provided",
    "List specific",
    "List recent",
)

#: The part of the map's list that can only be a placeholder, whoever wrote it:
#: fill-in marks, not English. Speaker notes are checked for these and for the
#: template's own notes text (`template_note_fragments`), and for nothing else.
#: The rest of the list ("Insert", "talk track", "List recent"...) is ordinary
#: wording that an engineer's own notes may well use ("Insert the new drive",
#: "my talk track for this slide"); on a slide those phrases mean an instruction
#: was left behind, in an engineer's notes they mean nothing.
NOTE_PLACEHOLDERS: tuple[str, ...] = ("XX", "ABC Company", "202x", "{", "<<")

#: A run of this many consecutive words from the template's notes is the template's
#: text, wherever it was pasted and however much of the line came with it. Shorter
#: runs are ordinary English ("no technical jargon"), and template lines shorter
#: than this are headings and stock phrases ("Rules:").
MIN_FRAGMENT_WORDS = 5

#: Sample content in the template that is not red and not caught by the list above.
TEMPLATE_SAMPLES: tuple[str, ...] = (
    "Name or Logo Here",
    "Overall IT Health Is Strong",
    "Security Posture Is Solid",
    "Proactive Management Is Reducing Noise",
    "Firewall and network security configurations",
    "New device installation, setup, and configuration",
    "Telephony (VOIP) configurations",
    "Other applicable reports",
)

#: The sample CX scores on slide 9. Real scores may equal them; then the caller says so.
SAMPLE_SCORES: tuple[str, ...] = ("88", "97")

DASHBOARD_SLIDE = 6
CX_SLIDE = 9


class RenderGateError(RuntimeError):
    """The deck still contains placeholder content. Nothing was written."""

    def __init__(self, problems: list[str]) -> None:
        self.problems = problems
        shown = "\n".join(f"  - {problem}" for problem in problems[:25])
        more = f"\n  ... and {len(problems) - 25} more" if len(problems) > 25 else ""
        super().__init__(f"render gate refused the deck ({len(problems)} problems):\n{shown}{more}")


@dataclass(slots=True)
class GateReport:
    problems: list[str] = field(default_factory=list)
    slides_checked: int = 0
    runs_checked: int = 0

    @property
    def passed(self) -> bool:
        return not self.problems


def _srgb(properties: Any) -> str | None:
    if properties is None:
        return None
    fill = properties.find(qn("a:solidFill"))
    if fill is None:
        return None
    colour = fill.find(qn("a:srgbClr"))
    return str(colour.get("val")).upper() if colour is not None else None


def effective_colour(run: Any) -> str | None:
    """The explicit colour governing a run: run, then paragraph default, then list style.

    Returns None when none of the three sets one, which means the text takes
    the theme's colour and is not red.
    """
    colour = _srgb(run.find(qn("a:rPr")))
    if colour:
        return colour
    paragraph = run.getparent()
    properties = paragraph.find(qn("a:pPr"))
    if properties is not None:
        colour = _srgb(properties.find(qn("a:defRPr")))
        if colour:
            return colour
    level = int(properties.get("lvl", "0")) if properties is not None else 0
    body = paragraph.getparent()
    styles = body.find(qn("a:lstStyle")) if body is not None else None
    if styles is not None:
        level_properties = styles.find(qn(f"a:lvl{level + 1}pPr"))
        if level_properties is not None:
            colour = _srgb(level_properties.find(qn("a:defRPr")))
            if colour:
                return colour
    return None


def red_text(slide: Any) -> list[str]:
    """Every piece of visible text on the slide whose effective colour is red."""
    found: list[str] = []
    for run in slide._element.iter(qn("a:r")):
        text_element = run.find(qn("a:t"))
        text = (text_element.text or "") if text_element is not None else ""
        if text.strip() and effective_colour(run) == RED:
            found.append(text.strip())
    return found


def all_text(slide: Any) -> list[str]:
    """Paragraph texts anywhere on the slide: shapes, groups, tables, chart titles excluded."""
    texts: list[str] = []
    for paragraph in slide._element.iter(qn("a:p")):
        text = "".join(t.text or "" for t in paragraph.iter(qn("a:t")))
        if text.strip():
            texts.append(text)
    return texts


def notes_text(slide: Any) -> str:
    """The slide's speaker notes, or an empty string. Never creates a notes page."""
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return (frame.text if frame is not None else "").strip()


def chart_text(slide: Any) -> list[str]:
    """Titles, series and category labels of the slide's charts: text the XML walk misses."""
    texts: list[str] = []
    for shape in slide.shapes:
        if not getattr(shape, "has_chart", False) or not shape.has_chart:
            continue
        chart = shape.chart
        if chart.has_title:
            texts.append(chart.chart_title.text_frame.text)
        for plot in chart.plots:
            texts.extend(str(category) for category in plot.categories)
            texts.extend(str(series.name) for series in plot.series)
    return texts


def _squash(text: str) -> str:
    return " ".join(text.split())


def _normalise(text: str) -> str:
    """Lower case, letters and digits only, single spaces: how notes are compared.

    Punctuation is dropped because the template's notes are full of curly quotes,
    bullets and non-breaking spaces that PowerPoint rewrites on save.
    """
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text.casefold()).split())


@lru_cache(maxsize=4)
def template_note_fragments(template: Path = TEMPLATE) -> tuple[str, ...]:
    """Every run of `MIN_FRAGMENT_WORDS` consecutive words in the template's notes.

    Read from the template itself, so the check follows the file when it changes.
    Runs do not cross lines. A note that contains any of them contains a piece of
    the template's notes, whether a whole note was left behind or half a sentence
    was pasted into an engineer's own text.
    """
    fragments: list[str] = []
    for slide in Presentation(str(template)).slides:
        for line in notes_text(slide).splitlines():
            words = _normalise(line).split()
            for start in range(len(words) - MIN_FRAGMENT_WORDS + 1):
                fragments.append(" ".join(words[start : start + MIN_FRAGMENT_WORDS]))
    return tuple(dict.fromkeys(fragments))


def note_problems(text: str, *, template: Path = TEMPLATE) -> list[str]:
    """Why this speaker-notes text cannot ship, or an empty list.

    Two things only: a fill-in mark from the slide map (`NOTE_PLACEHOLDERS`), or a
    line of the template's own notes. Used by the gate on the deck and by the
    review step on what an engineer types, so both refuse the same things.
    """
    problems: list[str] = []
    lowered = text.casefold()
    for mark in NOTE_PLACEHOLDERS:
        if mark in text if mark == "XX" else mark.casefold() in lowered:
            problems.append(f"placeholder mark {mark!r}")
    normalised = f" {_normalise(text)} "
    for fragment in template_note_fragments(template):
        if f" {fragment} " in normalised:
            # One is enough: a pasted sentence matches many overlapping runs.
            problems.append(f"the template's own notes text: {fragment!r}")
            break
    return problems


def _pictures(shapes: Any) -> list[Any]:
    found: list[Any] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            found.extend(_pictures(shape.shapes))
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            found.append(shape)
    return found


def check(
    slides: list[tuple[int, Any]],
    *,
    real_scores: frozenset[str] = frozenset(),
    allowed_notes: Mapping[int, str] | None = None,
    real_names: Callable[[str], list[str]] | None = None,
    private_names: Callable[[str], list[str]] | None = None,
    template: Path = TEMPLATE,
) -> GateReport:
    """Inspect the slides about to be saved. `slides` pairs a template slide number
    with the slide, because numbers shift once slides are dropped.

    Speaker notes are part of the deck: the template ships its own (presenter
    instructions and AI prompts on slides 1, 3, 6 and 7), so in every mode they are
    checked for that text and for fill-in marks (`note_problems`), and not for
    the ordinary phrases that are only a problem on a slide. With `allowed_notes`
    (a client or demo deck)
    a slide's notes must be exactly the engineer's talk track for it, or empty:
    no fact trace, no template instructions.

    `private_names` (any deck whose review has not turned contact names on) returns
    the customer's contact names found in a piece of text; a hit in a chart's labels
    is a problem. Slide text is not checked: an engineer may name a person in a
    sentence they wrote, and that is their call.

    `real_names` (a demo deck) returns the real names found in a piece of text;
    any hit, on a slide, in a chart or in the notes, is a problem.
    """
    report = GateReport()
    for number, slide in slides:
        report.slides_checked += 1
        where = f"slide {number}"
        for text in red_text(slide):
            report.problems.append(f"{where}: red placeholder text remains: {text[:70]!r}")
        texts = all_text(slide)
        report.runs_checked += len(texts)
        for text in texts:
            lowered = text.casefold()
            for forbidden in FORBIDDEN_STRINGS:
                hit = forbidden in text if forbidden == "XX" else forbidden.casefold() in lowered
                if hit:
                    report.problems.append(
                        f"{where}: placeholder string {forbidden!r} remains in {text[:70]!r}"
                    )
            for sample in TEMPLATE_SAMPLES:
                if sample.casefold() in lowered:
                    report.problems.append(f"{where}: template sample text remains: {sample!r}")
        notes = notes_text(slide)
        for problem in note_problems(notes, template=template):
            report.problems.append(f"{where}: speaker notes contain {problem}")
        if allowed_notes is not None and _squash(notes) != _squash(allowed_notes.get(number, "")):
            report.problems.append(
                f"{where}: speaker notes hold text that is not the engineer's talk track: "
                f"{notes[:60]!r}"
            )
        if real_names is not None:
            for text in (*texts, *chart_text(slide), notes):
                for name in real_names(text):
                    report.problems.append(
                        f"{where}: a real name survived in a demo deck: {name!r}"
                    )
        if private_names is not None:
            for text in chart_text(slide):
                for name in private_names(text):
                    report.problems.append(
                        f"{where}: a contact is named on a chart, and this review has not "
                        f"turned contact names on: {name!r}"
                    )
        if number == DASHBOARD_SLIDE and _pictures(slide.shapes):
            report.problems.append(
                f"{where}: the sample dashboard screenshot is still on the slide"
            )
        if number == CX_SLIDE:
            for text in texts:
                if text.strip() in SAMPLE_SCORES and text.strip() not in real_scores:
                    report.problems.append(
                        f"{where}: sample CX score {text.strip()!r} remains with no survey "
                        "result behind it"
                    )
    return report


def enforce(
    slides: list[tuple[int, Any]],
    *,
    real_scores: frozenset[str] = frozenset(),
    allowed_notes: Mapping[int, str] | None = None,
    real_names: Callable[[str], list[str]] | None = None,
    private_names: Callable[[str], list[str]] | None = None,
    template: Path = TEMPLATE,
) -> GateReport:
    report = check(
        slides,
        real_scores=real_scores,
        allowed_notes=allowed_notes,
        real_names=real_names,
        private_names=private_names,
        template=template,
    )
    if not report.passed:
        raise RenderGateError(report.problems)
    return report
