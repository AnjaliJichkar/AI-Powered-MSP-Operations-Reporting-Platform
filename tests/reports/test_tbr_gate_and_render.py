"""The render gate and the renderer, against the real 2026 template."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.reports.tbr import gate, review
from app.reports.tbr.config import QUOTE_PENDING
from app.reports.tbr.model import Item
from app.reports.tbr.render import TEMPLATE, render, slides_to_keep
from tests.reports.conftest import make_config, make_document, make_facts

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def run_in(xml: str) -> Any:
    root = etree.fromstring(xml.replace("<a:", "<a:").replace("xmlns:a=''", ""))
    return root.iter(f"{{{A}}}r").__next__()


def body(inner: str, list_style: str = "") -> str:
    return (
        f'<a:txBody xmlns:a="{A}"><a:bodyPr/><a:lstStyle>{list_style}</a:lstStyle>'
        f"{inner}</a:txBody>"
    )


RED = '<a:solidFill><a:srgbClr val="FF0000"/></a:solidFill>'
BLACK = '<a:solidFill><a:srgbClr val="000000"/></a:solidFill>'


# --- colour is resolved at all three levels -------------------------------------------------


def test_red_on_the_run() -> None:
    run = run_in(body(f"<a:p><a:r><a:rPr>{RED}</a:rPr><a:t>XX</a:t></a:r></a:p>"))
    assert gate.effective_colour(run) == "FF0000"


def test_red_as_the_paragraph_default() -> None:
    run = run_in(
        body(f"<a:p><a:pPr><a:defRPr>{RED}</a:defRPr></a:pPr><a:r><a:t>x</a:t></a:r></a:p>")
    )
    assert gate.effective_colour(run) == "FF0000"


def test_red_in_the_shapes_list_style_for_that_level() -> None:
    styles = f"<a:lvl2pPr><a:defRPr>{RED}</a:defRPr></a:lvl2pPr>"
    second = run_in(body('<a:p><a:pPr lvl="1"/><a:r><a:t>x</a:t></a:r></a:p>', styles))
    first = run_in(body("<a:p><a:r><a:t>x</a:t></a:r></a:p>", styles))
    assert gate.effective_colour(second) == "FF0000"
    assert gate.effective_colour(first) is None  # level one has no red default


def test_a_run_colour_overrides_a_red_default() -> None:
    run = run_in(
        body(
            f"<a:p><a:pPr><a:defRPr>{RED}</a:defRPr></a:pPr>"
            f"<a:r><a:rPr>{BLACK}</a:rPr><a:t>x</a:t></a:r></a:p>"
        )
    )
    assert gate.effective_colour(run) == "000000"


def test_the_template_really_uses_all_three_levels() -> None:
    """The map's claim, checked against the file: slide 5 run, slide 3 paragraph, slide 8 list."""
    slides = Presentation(str(TEMPLATE)).slides
    assert any("six" in text for text in gate.red_text(slides[4]))  # a red run
    assert any("technology trends" in text for text in gate.red_text(slides[2]))  # pPr default
    assert any("recent accomplishments" in text for text in gate.red_text(slides[7]))  # lstStyle


# --- the gate ---------------------------------------------------------------------------------


def test_the_untouched_template_fails_the_gate_on_every_placeholder_slide() -> None:
    presentation = Presentation(str(TEMPLATE))
    report = gate.check([(n, slide) for n, slide in enumerate(presentation.slides, start=1)])

    failing = {int(problem.split(":")[0].removeprefix("slide ")) for problem in report.problems}
    assert failing == {1, 2, 3, 5, 6, 7, 8, 9, 11, 12, 13, 14, 15, 17, 18, 19, 22}
    text = "\n".join(report.problems)
    assert "sample dashboard screenshot" in text
    assert "sample CX score '88'" in text
    assert "'<<'" in text and "'XX'" in text and "'202x'" in text
    assert "Overall IT Health Is Strong" in text  # sample content that is not red
    with pytest.raises(gate.RenderGateError):
        gate.enforce([(n, slide) for n, slide in enumerate(presentation.slides, start=1)])


def test_a_real_score_that_equals_a_sample_score_is_allowed() -> None:
    slide = Presentation(str(TEMPLATE)).slides[8]
    problems = gate.check([(9, slide)], real_scores=frozenset({"88", "97"})).problems
    assert not [p for p in problems if "sample CX score" in p]


# --- rendering ----------------------------------------------------------------------------------


async def test_a_review_renders_and_passes_the_gate(tmp_path: Path) -> None:
    document = await make_document()
    # The gap is single-source, so cybersecurity still has its measurements; trends need approval.
    result = render(document, tmp_path / "deck.pptx")

    assert result.gate.passed
    assert result.path.is_file()
    # Dropped: trends (unapproved research), CX (no survey), backup (deferred).
    assert set(result.dropped) == {3, 9, 15}
    deck = Presentation(str(result.path))
    assert len(deck.slides) == len(result.kept) == 20
    every_word = " ".join(
        gate.all_text(slide)[0] if gate.all_text(slide) else "" for slide in deck.slides
    )
    assert "Industry Trends" not in every_word  # the slide is gone, not blank


async def test_the_agenda_loses_the_lines_of_dropped_sections(tmp_path: Path) -> None:
    document = await make_document()
    render(document, tmp_path / "deck.pptx")
    agenda = [t.strip() for t in gate.all_text(Presentation(str(tmp_path / "deck.pptx")).slides[1])]

    assert "Industry Trends" not in agenda
    assert "Other applicable reports" not in agenda  # a sample line, never shipped
    assert {"Services Review", "Open Technical Items", "Budget Summary", "Reports"} <= set(agenda)


async def test_slide_six_gets_native_charts_and_loses_the_screenshot(tmp_path: Path) -> None:
    document = await make_document()
    render(document, tmp_path / "deck.pptx")
    deck = Presentation(str(tmp_path / "deck.pptx"))
    dashboard = next(s for s in deck.slides if "Customer Success" in " ".join(gate.all_text(s)))

    charts = [shape for shape in dashboard.shapes if getattr(shape, "has_chart", False)]
    assert len([c for c in charts if c.has_chart]) == 5
    assert not [s for s in dashboard.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    daily = next(c.chart for c in charts if c.has_chart and "Daily" in c.name)
    assert [s.name for s in daily.plots[0].series] == ["Tickets", "Completed"]


async def test_slide_twelve_is_retitled_and_priced_from_the_price_book(tmp_path: Path) -> None:
    document = await make_document()
    render(document, tmp_path / "deck.pptx")
    deck = Presentation(str(tmp_path / "deck.pptx"))
    texts = [" ".join(gate.all_text(slide)) for slide in deck.slides]

    replacement = next(t for t in texts if "Recommended for Replacement" in t)
    assert "ACME-WS01" in replacement
    assert "Hours: 3 hours. Cost: $1,450." in replacement  # from the price book, per device
    upgrade = next(t for t in texts if "Recommended for Upgrade" in t)
    assert f"Hours: {QUOTE_PENDING}. Cost: {QUOTE_PENDING}." in upgrade  # no entry: never a guess
    budget = next(t for t in texts if "Budget and Timeline" in t)
    assert "$900" in budget and QUOTE_PENDING in budget and "Q4 2026" in budget


async def test_no_price_book_means_every_figure_is_quote_pending(tmp_path: Path) -> None:
    config = make_config()
    config.price_book = {}
    document = await make_document(config=config)
    render(document, tmp_path / "deck.pptx")
    text = " ".join(
        " ".join(gate.all_text(s)) for s in Presentation(str(tmp_path / "deck.pptx")).slides
    )
    assert "$" not in text
    assert QUOTE_PENDING in text


async def test_a_section_with_no_facts_is_dropped_with_its_divider(tmp_path: Path) -> None:
    facts = make_facts(
        ninjaone__windows11__ineligible=0,
        ninjaone__windows11__eligible=0,
        ninjaone__disks__low=0,
        ninjaone__patches__failed=0,
        ninjaone__patches__critical_pending=0,
        ninjaone__antivirus__devices_unprotected=0,
        autotask__projects__current=0,
    )
    for table_id in list(facts.tables):
        if table_id.startswith(("ninjaone.", "autotask.projects.current")):
            del facts.tables[table_id]
    document = await make_document(facts=facts)

    keep = slides_to_keep(document)
    assert (
        not {10, 11, 12, 13, 14, 15} & keep
    )  # the whole Open Technical Items section, divider too
    render(document, tmp_path / "deck.pptx")
    agenda = [t.strip() for t in gate.all_text(Presentation(str(tmp_path / "deck.pptx")).slides[1])]
    assert "Open Technical Items" not in agenda


async def test_placeholder_text_on_a_slide_stops_the_render_and_writes_nothing(
    tmp_path: Path,
) -> None:
    document = await make_document()
    # Past the review step's own check, straight into the document: the gate is the last line.
    document.section("accomplishments").engineer_items.append(
        Item(text="Migrated XX mailboxes in July 202x", cites=["engineer:x"], author="engineer")
    )
    target = tmp_path / "deck.pptx"

    with pytest.raises(gate.RenderGateError) as refused:
        render(document, target)

    assert not target.exists()
    assert any("'XX'" in problem for problem in refused.value.problems)
    assert any("'202x'" in problem for problem in refused.value.problems)


async def test_engineer_input_puts_a_slide_back_and_approval_gates_the_web_slide(
    tmp_path: Path,
) -> None:
    document = await make_document()
    review.set_scores(document, nps="71", csat="97", pqual=None, comments=["Quick"], by="sam")
    trends = document.section("trends")
    trends.items = [
        Item(
            text="Insurers now ask for proof of controls.",
            cites=["web"],
            author="model",
            url="https://www.cisa.gov/insurance",
        )
    ]
    assert 3 not in slides_to_keep(document)  # researched, not approved: stays off the deck
    review.approve(document, "trends", by="sam")

    result = render(document, tmp_path / "deck.pptx")

    assert {3, 9} <= set(result.kept)
    deck = Presentation(str(result.path))
    cx = next(s for s in deck.slides if "CX Metrics" in " ".join(gate.all_text(s)))
    shown = {t.strip() for t in gate.all_text(cx)}
    assert {"71", "97"} <= shown  # 97 is real here, and the gate was told so
    assert not {"88"} & shown
    assert "(PQUAL)" not in " ".join(shown)  # no score, so no label either
    trends_slide = next(s for s in deck.slides if gate.all_text(s)[:1] == ["Industry Trends"])
    assert "(Source: cisa.gov)" in " ".join(gate.all_text(trends_slide))
    assert gate.notes_text(trends_slide) == ""  # a client deck: no trace in the notes
