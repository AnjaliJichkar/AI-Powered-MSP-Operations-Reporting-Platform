"""Contact privacy: people on the tickets-by-contact chart are Contact A, B unless turned on."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from app.reports.tbr import gate, review
from app.reports.tbr.model import ChartData, ChartSeries, ReviewDocument
from app.reports.tbr.render import (
    CONTACT_CHART,
    anonymised_contacts,
    contact_names_in,
    render,
)
from tests.reports.conftest import make_config, make_document, make_facts

PEOPLE = ["Dana Whitfield", "Omar Reyes", "Priya Nair"]


def facts_with_people() -> Any:
    """Three named people, an unnamed contact and the automation bar; bottom-up like a bar chart."""
    facts = make_facts()
    facts.charts[CONTACT_CHART] = ChartData(
        id=CONTACT_CHART,
        title="Who raised tickets",
        kind="bar",
        categories=[
            "Automation and system accounts",
            "Unnamed contact",
            "Priya Nair",
            "Omar Reyes",
            "Dana Whitfield",  # last is drawn at the top: the most tickets
        ],
        series=[ChartSeries(name="Tickets", values=[20.0, 2.0, 4.0, 9.0, 31.0])],
        definition="Test chart.",
        source="autotask",
    )
    facts.names["contacts"] = [*PEOPLE, "Someone Offchart"]
    return facts


def contact_chart(result: Any) -> tuple[list[str], list[float]]:
    deck = Presentation(str(result.path))
    slide = list(deck.slides)[result.kept.index(6)]
    chart = next(
        s.chart
        for s in slide.shapes
        if getattr(s, "has_chart", False) and s.has_chart and s.name == "Chart: Who raised tickets"
    )
    plot = chart.plots[0]
    return [str(c) for c in plot.categories], list(plot.series[0].values)


def test_labels_follow_rank_and_bars_that_are_not_people_keep_theirs() -> None:
    real = facts_with_people().charts[CONTACT_CHART]

    hidden = anonymised_contacts(real)

    assert hidden.categories == [
        "Automation and system accounts",
        "Unnamed contact",
        "Contact C",
        "Contact B",
        "Contact A",  # the most tickets
    ]
    assert hidden.series == real.series  # every figure, in the same order
    assert real.categories[-1] == "Dana Whitfield"  # the stored chart is not touched


async def test_a_deck_hides_names_by_default_in_client_and_internal_modes(tmp_path: Path) -> None:
    document = await make_document(facts=facts_with_people())
    # No override, and the franchisee default is off: people are hidden.
    assert document.show_contact_names is None and document.people_names_shown is False

    for mode in ("client", "internal"):
        result = render(document, tmp_path / f"{mode}.pptx", mode=mode)  # type: ignore[arg-type]
        labels, values = contact_chart(result)
        assert labels[-3:] == ["Contact C", "Contact B", "Contact A"]
        assert values == [20.0, 2.0, 4.0, 9.0, 31.0]
        deck = Presentation(str(result.path))
        charted = " ".join(t for slide in deck.slides for t in gate.chart_text(slide))
        assert not [name for name in PEOPLE if name in charted]
    assert document.facts.charts[CONTACT_CHART].categories[-1] == "Dana Whitfield"


async def test_an_engineer_turns_names_on_for_one_review_and_it_is_recorded(
    tmp_path: Path,
) -> None:
    document = await make_document(facts=facts_with_people())

    review.set_contact_names(document, True, by="sam")
    labels, _ = contact_chart(render(document, tmp_path / "named.pptx"))

    assert labels[-3:] == ["Priya Nair", "Omar Reyes", "Dana Whitfield"]
    assert (document.show_contact_names, document.show_contact_names_by) == (True, "sam")
    other = await make_document(facts=facts_with_people())
    assert other.show_contact_names is None  # per review, not a standing setting

    # Hiding is an override too, and who did it is recorded either way.
    review.set_contact_names(document, False, by="alex")
    labels, _ = contact_chart(render(document, tmp_path / "hidden.pptx"))
    assert labels[-1] == "Contact A" and document.show_contact_names_by == "alex"
    assert document.people_names_shown is False

    # Clearing the override goes back to the franchisee's default.
    review.set_contact_names(document, None, by="alex")
    assert (document.show_contact_names, document.show_contact_names_by) == (None, None)
    assert document.people_names_shown is document.people_names_default


async def test_demo_mode_is_unaffected_by_the_setting(tmp_path: Path) -> None:
    hidden = await make_document(facts=facts_with_people())
    named = await make_document(facts=facts_with_people())
    review.set_contact_names(named, True, by="sam")

    first = contact_chart(render(hidden, tmp_path / "a.pptx", mode="demo"))
    second = contact_chart(render(named, tmp_path / "b.pptx", mode="demo"))

    assert first == second  # a demo deck has its own pseudonyms for everyone, either way
    assert not [name for name in PEOPLE if name in " ".join(first[0])]


async def test_the_gate_refuses_a_named_chart_when_names_are_off(tmp_path: Path) -> None:
    document = await make_document(facts=facts_with_people())
    review.set_contact_names(document, True, by="sam")
    result = render(document, tmp_path / "named.pptx")  # a deck with real names on its chart
    deck = Presentation(str(result.path))
    slides = list(zip(result.kept, deck.slides, strict=True))
    checker = contact_names_in(document)
    assert "Someone Offchart" in checker.exact  # every contact seen, not only the charted

    report = gate.check(slides, private_names=checker.survivors)

    named = [p for p in report.problems if "has not turned contact names on" in p]
    assert len(named) == 3 and all(p.startswith("slide 6:") for p in named)
    assert gate.check(slides).passed  # with names on, the same deck is fine


# --- the franchisee default, and the review's override of it ----------------------------------


async def test_the_franchisee_default_names_people_on_client_and_internal_decks(
    tmp_path: Path,
) -> None:
    """`show_people_names` on the franchisee config is the default for every review;
    nothing per review needs setting, and who is named is recorded as the default."""
    document = await make_document(
        facts=facts_with_people(), config=make_config(show_people_names=True)
    )
    assert document.show_contact_names is None  # no override...
    assert document.people_names_default is True and document.people_names_shown is True

    for mode in ("client", "internal"):
        labels, _ = contact_chart(render(document, tmp_path / f"{mode}.pptx", mode=mode))  # type: ignore[arg-type]
        assert labels[-3:] == ["Priya Nair", "Omar Reyes", "Dana Whitfield"]

    # ...and an engineer can still hide them for this one review.
    review.set_contact_names(document, False, by="sam")
    labels, _ = contact_chart(render(document, tmp_path / "hidden.pptx"))
    assert labels[-1] == "Contact A" and document.people_names_shown is False


async def test_a_demo_deck_pseudonymises_people_whatever_the_franchisee_says(
    tmp_path: Path,
) -> None:
    document = await make_document(
        facts=facts_with_people(), config=make_config(show_people_names=True)
    )
    labels, _ = contact_chart(render(document, tmp_path / "demo.pptx", mode="demo"))
    assert not [name for name in PEOPLE if name in " ".join(labels)]
    assert labels[-1].startswith("Contact ")


async def test_the_client_gate_follows_the_setting(tmp_path: Path) -> None:
    """With names on, a named chart passes the client render; with them off, the same
    review's client render prints Contact A and a named chart is refused."""
    document = await make_document(
        facts=facts_with_people(), config=make_config(show_people_names=True)
    )
    named = render(document, tmp_path / "named.pptx")  # passes: the gate is not armed
    assert contact_chart(named)[0][-1] == "Dana Whitfield"

    document.franchisee["show_people_names"] = False  # what set_people_names writes back
    hidden = render(document, tmp_path / "hidden.pptx")
    assert contact_chart(hidden)[0][-1] == "Contact A"
    deck = Presentation(str(named.path))
    slides = list(zip(named.kept, deck.slides, strict=True))
    report = gate.check(slides, private_names=contact_names_in(document).survivors)
    assert any("has not turned contact names on" in p for p in report.problems)


async def test_a_stored_review_from_before_the_default_carries_no_override() -> None:
    """Old rows hold `show_contact_names: false` with nobody named, which was the old
    default and not a decision: it reads as no override, so the franchisee decides."""
    document = await make_document(facts=facts_with_people())
    data = document.model_dump(mode="json")
    data["show_contact_names"], data["show_contact_names_by"] = False, None
    data["franchisee"]["show_people_names"] = True
    loaded = ReviewDocument.model_validate(data)
    assert loaded.show_contact_names is None and loaded.people_names_shown is True

    data["show_contact_names"], data["show_contact_names_by"] = False, "alex"  # a decision
    assert ReviewDocument.model_validate(data).people_names_shown is False


@pytest.mark.parametrize(
    ("stored", "by", "override", "shown_when_default_on"),
    [
        (True, "sam", True, True),
        (False, "alex", False, False),
        (None, None, None, True),
        ("missing", None, None, True),  # a document from before the field existed
        (False, None, None, True),  # the old default, not a decision
    ],
)
async def test_a_stored_override_loads_in_every_shape_it_can_be_stored_in(
    stored: Any, by: str | None, override: bool | None, shown_when_default_on: bool
) -> None:
    document = await make_document(facts=facts_with_people())
    data = document.model_dump(mode="json")
    data["franchisee"]["show_people_names"] = True
    if stored == "missing":
        del data["show_contact_names"]
    else:
        data["show_contact_names"] = stored
    data["show_contact_names_by"] = by

    loaded = ReviewDocument.model_validate(data)
    assert loaded.show_contact_names is override
    assert loaded.people_names_shown is shown_when_default_on
    # What it writes back loads again, unchanged.
    again = ReviewDocument.model_validate(loaded.model_dump(mode="json"))
    assert again.show_contact_names is override and again.show_contact_names_by == (
        by if override is not None else None
    )
