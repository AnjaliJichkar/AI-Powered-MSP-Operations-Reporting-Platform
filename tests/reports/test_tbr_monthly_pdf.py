"""The Monthly Report PDF: the review's second output, from the same stored facts.

No vendor, no model, no database: the synthetic review from `conftest.py` with the
facts the PDF adds (lifecycle tables, third-party patch installs, the first-response
share), rendered with ReportLab and read back with pypdf.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pypdf import PdfReader

from app.reports.tbr import review
from app.reports.tbr.collect import ROLE_DISAGREEMENTS
from app.reports.tbr.demo import (
    STAMP,
    build_pseudonyms,
    real_name_checker,
    real_name_terms,
    substring_survivors,
)
from app.reports.tbr.gate import RenderGateError
from app.reports.tbr.model import ChartData, ChartSeries
from app.reports.tbr.monthly import (
    LIFECYCLE_COLUMNS,
    LIFECYCLE_SERVERS,
    LIFECYCLE_WORKSTATIONS,
    NOT_REPORTED,
    free_band,
)
from app.reports.tbr.pdf import (
    MONTHLY_SECTION,
    Pseudonyms,
    _age_band,
    default_pdf_output,
    gate_texts,
    render_pdf,
)
from app.reports.tbr.render import CHART_LAYOUT, CONTACT_CHART, render, slides_to_keep
from tests.reports.conftest import NOW, fact, make_config, make_document, make_facts, table

PEOPLE = ["Dana Whitfield", "Omar Reyes"]
LOGINS = ["ACME\\dwhitfield", "ACME\\oreyes", "dana.whitfield@acme.test"]
ADMIN = "ACME\\admin"
# NinjaOne locations are free text: one embeds the site's name, one a person's.
LOCATIONS = ["Acme Legal - Main office", "Dana Whitfield's office", "Server room"]


def monthly_facts() -> Any:
    """The synthetic facts plus everything only the PDF prints."""
    # The device counts are the lifecycle tables' populations: four workstations,
    # one physical server, no virtual one, as the collector would have counted them.
    facts = make_facts(
        ninjaone__devices__workstations=4,
        ninjaone__devices__servers_physical=1,
        ninjaone__devices__servers_virtual=0,
    )
    facts.add(
        fact("autotask.sla.first_response_met_percent", 91.3, source="autotask", display="91.3%")
    )
    facts.add(fact("autotask.sla.first_response_measured", 110, source="autotask"))
    facts.add(fact("ninjaone.patches.attempted", 324))
    facts.add(fact("ninjaone.patches.top_failed_device", "ACME-WS02"))
    facts.add(fact("ninjaone.software_patches.installed", 88))
    facts.add(fact("ninjaone.software_patches.failed", 3))
    facts.add(fact("ninjaone.software_patches.attempted", 91))
    facts.add(fact("ninjaone.software_patches.success_percent", 96.7, display="96.7%"))
    facts.add(fact("ninjaone.software_patches.failed_devices", 2))
    facts.add(fact("ninjaone.software_patches.top_failed_device", "ACME-WS01"))
    facts.tables["ninjaone.patches.failed.devices"] = table(
        "ninjaone.patches.failed.devices", ["Device", "Failed"], [["ACME-WS02", "9"]]
    )
    facts.charts[CONTACT_CHART] = ChartData(
        id=CONTACT_CHART,
        title="Who raised tickets",
        kind="bar",
        categories=["Automation and system accounts", "Omar Reyes", "Dana Whitfield"],
        series=[ChartSeries(name="Tickets", values=[20.0, 9.0, 31.0])],
        definition="Test chart.",
        source="autotask",
    )
    facts.names["contacts"] = list(PEOPLE)
    # device, user, make, model, os, ram, cpu, location, used, total, free, free %, age
    rows = [
        ["ACME-WS01", LOGINS[0], "Dell Inc.", "OptiPlex 7090", "Windows 11 Pro", "16.0 GB",
         "Intel(R) Core(TM) i5-10500", LOCATIONS[0], "214.6 GB", "477.0 GB", "262.3 GB",
         "55%", "6"],
        ["ACME-WS02", LOGINS[1], "Dell Inc.", "OptiPlex 7040", "Windows 10 Pro", "8.0 GB",
         "Intel(R) Core(TM) i5-6500", LOCATIONS[0], "333.9 GB", "477.0 GB", "143.1 GB",
         "30%", "10"],
        ["ACME-WS03", LOGINS[2], "HP", "EliteDesk 800", "Windows 11 Pro", "16.0 GB",
         "Intel(R) Core(TM) i7-12700", LOCATIONS[1], "419.8 GB", "477.0 GB", "57.2 GB", "12%", "4"],
        ["ACME-WS04", NOT_REPORTED, "HP", "EliteDesk 800", "Windows 11 Pro", NOT_REPORTED,
         "Unknown CPU", LOCATIONS[1], NOT_REPORTED, NOT_REPORTED, NOT_REPORTED, NOT_REPORTED, ""],
        ["ACME-DC01", ADMIN, "Dell Inc.", "PowerEdge R740", "Windows Server 2019", "64.0 GB",
         "Intel(R) Xeon(R) Silver 4210", LOCATIONS[2], "1.4 TB", "1.8 TB", "0.4 TB", "22%", "7"],
    ]  # fmt: skip
    facts.tables[LIFECYCLE_WORKSTATIONS] = table(
        LIFECYCLE_WORKSTATIONS, list(LIFECYCLE_COLUMNS), rows[:-1]
    )
    facts.tables[LIFECYCLE_SERVERS] = table(LIFECYCLE_SERVERS, list(LIFECYCLE_COLUMNS), rows[-1:])
    facts.names["users"] = [*LOGINS, ADMIN]
    facts.names["locations"] = list(LOCATIONS)
    return facts


async def monthly_document(**kwargs: Any) -> Any:
    return await make_document(facts=kwargs.pop("facts", None) or monthly_facts(), **kwargs)


def pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def deck_words(result: Any) -> str:
    from app.reports.tbr import gate

    deck = Presentation(str(result.path))
    return "\n".join(
        " \n ".join([*gate.all_text(s), *gate.chart_text(s), gate.notes_text(s)])
        for s in deck.slides
    )


# --- the approval step -------------------------------------------------------------------------


async def test_the_review_carries_the_monthly_report_as_a_section_off_the_deck() -> None:
    document = await monthly_document()
    section = document.section(MONTHLY_SECTION)

    assert section.slide == 0 and section.requires_approval and section.approved_by is None
    assert section.status == "filled" and section.narrative == "code"
    assert [item.author for item in section.items] == ["code"] * len(section.items)
    assert 0 not in slides_to_keep(document)  # not a slide: never on the deck
    line = next(row for row in review.status_report(document) if row["slide"] == 0)
    assert line["on_deck"] == "no" and line["status"] == "awaiting_engineer"

    # Approval is the ordinary review action, and any edit to the review withdraws it.
    review.approve(document, MONTHLY_SECTION, by="sam")
    assert document.section(MONTHLY_SECTION).approved_by == "sam"


async def test_a_client_pdf_needs_approval_and_the_other_modes_say_they_are_drafts(
    tmp_path: Path,
) -> None:
    document = await monthly_document()

    with pytest.raises(RenderGateError, match="not been approved"):
        render_pdf(document, tmp_path / "client.pdf", now=NOW)
    assert not (tmp_path / "client.pdf").exists()

    internal = render_pdf(document, tmp_path / "internal.pdf", mode="internal", now=NOW)
    assert "Draft: not approved" in pdf_text(internal.path)
    assert "facts: autotask.chart.daily" in pdf_text(internal.path)  # the trace, internal only

    review.approve(document, MONTHLY_SECTION, by="sam")
    client = render_pdf(document, tmp_path / "client.pdf", now=NOW)
    text = pdf_text(client.path)
    assert "Draft" not in text and "facts:" not in text and STAMP not in text
    assert client.sections == ["Ticketing", "Monitoring", "Backup"] and client.omitted == []
    assert client.pages >= 6


async def test_a_review_that_predates_the_pdf_is_told_to_generate_again(tmp_path: Path) -> None:
    document = await monthly_document()
    document.sections = [s for s in document.sections if s.key != MONTHLY_SECTION]
    with pytest.raises(RenderGateError, match="generate it again"):
        render_pdf(document, tmp_path / "client.pdf", now=NOW)


# --- what is printed -------------------------------------------------------------------------


async def test_the_pdf_prints_the_stored_facts_and_no_sentence_of_its_own(
    tmp_path: Path,
) -> None:
    document = await monthly_document()
    review.approve(document, MONTHLY_SECTION, by="sam")
    result = render_pdf(document, tmp_path / "client.pdf", now=NOW)
    text = pdf_text(result.path)

    # Every figure drawn is a stored fact's display string, by id, and is on the page.
    assert result.figures
    for fact_id, display in result.figures.items():
        assert document.facts.facts[fact_id].display == display
        assert display in text, fact_id
    # The totals under the daily chart are the created and completed facts.
    assert result.figures["autotask.tickets.created"] == "120"
    assert result.figures["autotask.tickets.completed"] == "115"
    # Both gauges, with the target and the written definition.
    assert "91.3%" in text and "82%" in text and "Target: 90%" in text
    assert "Out of 110 tickets" in text and "Out of 115 tickets" in text
    # Automated contacts are a separate count, not bars.
    assert "shown separately: 20" in text
    assert "Automation and system accounts" not in text
    # Both patch cards.
    assert "Status of OS patches" in text and "Status of third-party patches" in text
    assert "95.7%" in text and "96.7%" in text
    # The charts the deck also draws were drawn from the same data.
    assert set(result.charts) >= {"autotask.chart.daily", "autotask.chart.issue_type"}


async def test_missing_storage_reads_not_reported_and_never_zero(tmp_path: Path) -> None:
    document = await monthly_document()
    review.approve(document, MONTHLY_SECTION, by="sam")
    text = pdf_text(render_pdf(document, tmp_path / "client.pdf", now=NOW).path)

    # The last workstation row runs from its name to the table's footnote.
    row = " ".join(text.split("ACME-WS04", 1)[1].split("Test table.", 1)[0].split())
    assert row.count(NOT_REPORTED) >= 5 and "0.0 GB" not in row and " 0%" not in row
    assert "Estimated age" in text and "Storage" in text  # the legend
    assert "Test table." in text  # the table's definition: where the age method is footnoted


def test_storage_bands_and_age_bands() -> None:
    assert free_band(55) == "healthy" and free_band(40) == "healthy"
    assert free_band(39.9) == "low" and free_band(20) == "low"
    assert free_band(19.9) == "critical" and free_band(0) == "critical"
    assert free_band(None) == "unknown"
    assert _age_band("3") == "healthy" and _age_band("4") == "low" and _age_band("6") == "low"
    assert _age_band("7") == "critical" and _age_band("") == "unknown"


async def test_sub_issue_types_appear_when_other_or_unset_is_more_than_half(
    tmp_path: Path,
) -> None:
    facts = monthly_facts()
    facts.add(
        fact("autotask.tickets.issue_type_unset_share", 62.5, source="autotask", display="62.5%")
    )
    document = await monthly_document(facts=facts)
    review.approve(document, MONTHLY_SECTION, by="sam")
    text = pdf_text(render_pdf(document, tmp_path / "a.pdf", now=NOW).path)
    assert "largely unset" in text or "no breakdown is shown" in text

    facts.charts["autotask.chart.sub_issue_type"] = ChartData(
        id="autotask.chart.sub_issue_type",
        title="Tickets by sub-issue type",
        kind="column",
        categories=["Outlook", "Teams"],
        series=[ChartSeries(name="Tickets", values=[40.0, 35.0])],
        definition="Test chart.",
        source="autotask",
    )
    document = await monthly_document(facts=facts)
    review.approve(document, MONTHLY_SECTION, by="sam")
    text = pdf_text(render_pdf(document, tmp_path / "b.pdf", now=NOW).path)
    assert "Sub-issue types (62.5% of tickets" in text and "Outlook" in text


async def test_a_review_with_no_ticket_facts_omits_ticketing_and_says_so(
    tmp_path: Path,
) -> None:
    facts = monthly_facts()
    for fact_id in [f for f in facts.facts if f.startswith("autotask.")]:
        del facts.facts[fact_id]
    for chart_id in [c for c in facts.charts if c.startswith("autotask.")]:
        del facts.charts[chart_id]
    document = await monthly_document(facts=facts)
    review.approve(document, MONTHLY_SECTION, by="sam")

    result = render_pdf(document, tmp_path / "client.pdf", now=NOW)
    assert result.sections == ["Monitoring", "Backup"]
    assert ("Ticketing", "the review holds no Autotask ticket facts") in result.omitted
    assert "Ticketing" not in pdf_text(result.path)


async def test_backup_says_so_in_one_line_when_cove_is_not_connected(tmp_path: Path) -> None:
    """A site with no Cove customer gets a sentence, not a page of zeros.

    It still names what NinjaOne holds: "not assessed" on its own is a shrug,
    and beside a count it is a size. This test used to assert the empty shape
    of every card the section would one day fill; now that Cove fills them,
    what is worth pinning is the line a site without it gets.
    """
    document = await monthly_document()
    review.approve(document, MONTHLY_SECTION, by="sam")
    text = pdf_text(render_pdf(document, tmp_path / "client.pdf", now=NOW).path)
    flat = " ".join(text.split())
    assert "Backup coverage, status and storage" in flat  # the subtitle
    assert "Not yet assessed" in flat and "N-able Cove" in flat
    assert "4 workstations" in flat  # the NinjaOne side, so the gap has a size
    # And none of the Cove figures are invented in its absence.
    for absent in ("protected:", "Backup success", "Not backed up by Cove"):
        assert absent not in flat, absent


# --- privacy and the gate ------------------------------------------------------------------


async def test_people_are_hidden_unless_names_are_on_and_a_demo_names_no_one(
    tmp_path: Path,
) -> None:
    document = await monthly_document()
    review.approve(document, MONTHLY_SECTION, by="sam")

    private = pdf_text(render_pdf(document, tmp_path / "client.pdf", now=NOW).path)
    for name in (*PEOPLE, *LOGINS, ADMIN):
        assert name not in private, name
    assert "Contact A" in private and "User A" in private
    # A location named after a person is withheld too; the site's own name is not private.
    flat = " ".join(private.split())
    assert "Location withheld" in flat and LOCATIONS[0] in flat

    review.set_contact_names(document, True, by="sam")
    named = pdf_text(render_pdf(document, tmp_path / "named.pdf", now=NOW).path)
    assert "Dana Whitfield" in named and LOGINS[0] in named

    demo = pdf_text(render_pdf(document, tmp_path / "demo.pdf", mode="demo", now=NOW).path)
    names = build_pseudonyms(document)
    for real in (*PEOPLE, *LOGINS, ADMIN, "ACME-WS01", "Acme Legal", "Example IT Partners"):
        assert real in names.exact or real in names.loose
        assert real not in demo, real
    assert STAMP in demo and "DEVICE-001" in demo and "User A" in demo


def test_the_gate_refuses_tokens_placeholders_and_names_that_must_not_appear() -> None:
    private = Pseudonyms(exact={"Dana Whitfield": ""})
    with pytest.raises(RenderGateError, match="fact token"):
        gate_texts(["{{autotask.tickets.created}} tickets"], private_names=None, real_names=None)
    with pytest.raises(RenderGateError, match="placeholder"):
        gate_texts(["ABC Company"], private_names=None, real_names=None)
    with pytest.raises(RenderGateError, match="private"):
        gate_texts(["Dana Whitfield raised 31"], private_names=private, real_names=None)
    demo_check = real_name_checker(Pseudonyms(exact={"ACME-WS01": ""}), ["Acme Legal"])
    with pytest.raises(RenderGateError, match="demo"):
        gate_texts(["ACME-WS01"], private_names=None, real_names=demo_check)
    with pytest.raises(RenderGateError, match="Acme Legal"):
        gate_texts(["Office: ACME LEGAL-HOUSTON"], private_names=None, real_names=demo_check)
    gate_texts(["Contact A raised 31 tickets"], private_names=private, real_names=None)


# --- the demo gate: every real name, as a substring, anywhere ------------------------------


async def test_the_demo_check_list_holds_every_name_and_matches_substrings() -> None:
    document = await monthly_document()
    terms = real_name_terms(document)

    for expected in (
        "Acme Legal",
        "acme-legal",
        "Example IT Partners",
        "example-it",
        *PEOPLE,
        *LOGINS,
        ADMIN,
        "ACME-WS01",
        "ACME-DC01",
        *LOCATIONS,
    ):
        assert expected in terms, expected
    assert terms == sorted(terms, key=lambda t: (-len(t), t))  # longest first
    assert all(len(t) >= 3 for t in terms)

    # Case-insensitive, inside other words, inside a label the mapping never saw.
    assert substring_survivors("Ship to ACME LEGAL-HOUSTON on Monday", terms) == ["Acme Legal"]
    assert substring_survivors("dana whitfield's desk", terms) == ["Dana Whitfield"]
    assert substring_survivors("The front desk printer", terms) == []
    # The price of a substring rule: a device called RECEPTION makes the word a hit.
    # The gate is stricter than the replacement on purpose; the replacement stays
    # whole-token so the word itself is never rewritten.
    assert substring_survivors("The reception printer", terms) == ["RECEPTION"]
    assert substring_survivors("acme-ws01 rebuilt", terms) == ["ACME-WS01"]


async def test_a_demo_render_replaces_locations_and_the_gate_passes_the_substring_check(
    tmp_path: Path,
) -> None:
    document = await monthly_document()
    review.approve(document, MONTHLY_SECTION, by="sam")
    names = build_pseudonyms(document)
    # Each location maps whole, and before the site name inside it would.
    assert names.loose[LOCATIONS[0]] == "Location A"
    assert names.loose[LOCATIONS[1]] == "Location B"
    assert names.loose[LOCATIONS[2]] == "Location C"
    assert names.apply(LOCATIONS[0]) == "Location A"
    assert names.apply(LOCATIONS[1]) == "Location B"

    result = render_pdf(document, tmp_path / "demo.pdf", mode="demo", now=NOW)
    text = pdf_text(result.path)
    terms = real_name_terms(document)
    assert substring_survivors(text, terms) == []
    assert "Location A" in text and "Location B" in text and "Location C" in text
    for location in LOCATIONS:
        assert location.casefold() not in text.casefold()

    # The deck's demo render runs the same check on every slide, chart and note.
    deck = render(document, tmp_path / "demo.pptx", mode="demo")
    assert substring_survivors(deck_words(deck), terms) == []


async def test_a_real_name_hidden_inside_free_text_is_refused_by_the_demo_gate(
    tmp_path: Path,
) -> None:
    """A vendor free-text field the mapping never covers (a model string here)
    carries the site name run into a digit, so it is not a whole token: the
    mapping misses it, the substring check does not, and no file is written."""
    facts = monthly_facts()
    table = facts.tables[LIFECYCLE_WORKSTATIONS]
    column = table.columns.index("Model")
    table.rows[0][column] = "Built for Acme Legal2"
    document = await monthly_document(facts=facts)
    review.approve(document, MONTHLY_SECTION, by="sam")

    with pytest.raises(RenderGateError, match="Acme Legal"):
        render_pdf(document, tmp_path / "demo.pdf", mode="demo", now=NOW)
    assert not (tmp_path / "demo.pdf").exists()


async def test_the_file_name_says_which_pdf_it_is(tmp_path: Path) -> None:
    document = await monthly_document()
    assert default_pdf_output(document, tmp_path, now=NOW).name == (
        "acme-legal-monthly-2026-08-20260921-090000.pdf"
    )
    assert default_pdf_output(document, tmp_path, now=NOW, mode="internal").name.endswith(
        "-internal.pdf"
    )
    demo = default_pdf_output(document, tmp_path, now=NOW, mode="demo").name
    assert "acme" not in demo and demo.startswith("sample-client-monthly-")


# --- agreement with the deck -------------------------------------------------------------------


def _deck_charts(result: Any) -> dict[str, dict[str, dict[str, float]]]:
    """Every native chart on the deck, title -> series -> category -> value."""
    found: dict[str, dict[str, dict[str, float]]] = {}
    for slide in Presentation(str(result.path)).slides:
        for shape in slide.shapes:
            if not (getattr(shape, "has_chart", False) and shape.has_chart):
                continue
            plot = shape.chart.plots[0]
            categories = [str(c) for c in plot.categories]
            found[shape.name.removeprefix("Chart: ")] = {
                series.name: dict(zip(categories, series.values, strict=True))
                for series in plot.series
            }
    return found


async def test_every_number_shared_by_the_pdf_and_the_deck_is_identical(tmp_path: Path) -> None:
    """One stored review, both outputs, client mode: a fact the PDF draws and the deck
    cites prints the same display string on both; a chart both draw carries the same
    value for every category."""
    document = await monthly_document()
    review.approve(document, MONTHLY_SECTION, by="sam")

    deck = render(document, tmp_path / "client.pptx")
    pdf = render_pdf(document, tmp_path / "client.pdf", now=NOW)
    on_deck = deck_words(deck)
    on_page = pdf_text(pdf.path)

    cited = {
        cite
        for section in document.sections
        if section.slide in deck.kept
        for item in section.all_items()
        for cite in item.cites
    }
    shared = sorted(set(pdf.figures) & cited)
    assert len(shared) >= 4, shared
    for fact_id in shared:
        display = document.facts.facts[fact_id].display
        assert pdf.figures[fact_id] == display
        pattern = rf"(?<![\d.]){re.escape(display)}(?![\d.])"
        assert re.search(pattern, on_page), (fact_id, "pdf")
        assert re.search(pattern, on_deck), (fact_id, "deck")

    deck_charts = _deck_charts(deck)
    compared = 0
    for chart_id, _ in CHART_LAYOUT:
        if chart_id not in pdf.charts:
            continue
        title = document.facts.charts[chart_id].title
        assert title in deck_charts, chart_id
        for series_name, values in pdf.charts[chart_id].items():
            for category, value in values.items():
                assert deck_charts[title][series_name][category] == value, (chart_id, category)
                compared += 1
    assert compared >= 40  # the daily chart alone is 31 days of two series

    # The lifecycle tables and the deck's device counts are one population, decided
    # by one function (`hardware.device_role`): the rows on the page are the counts.
    workstations = document.facts.tables[LIFECYCLE_WORKSTATIONS]
    servers = document.facts.tables[LIFECYCLE_SERVERS]
    assert len(workstations.rows) == document.facts.number("ninjaone.devices.workstations")
    assert len(servers.rows) == (
        document.facts.number("ninjaone.devices.servers_physical") or 0
    ) + (document.facts.number("ninjaone.devices.servers_virtual") or 0)
    for fact_id in (
        "ninjaone.devices.workstations",
        "ninjaone.devices.servers_physical",
        "ninjaone.devices.servers_virtual",
    ):
        display = document.facts.facts[fact_id].display
        assert re.search(rf"(?<![\d.]){re.escape(display)}(?![\d.])", on_deck), fact_id
    assert f"{len(workstations.rows)} workstations" in on_page or len(workstations.rows) == 0


async def test_no_servers_is_a_sentence_not_an_empty_table(tmp_path: Path) -> None:
    facts = monthly_facts()
    facts.tables[LIFECYCLE_SERVERS] = table(LIFECYCLE_SERVERS, list(LIFECYCLE_COLUMNS), [])
    document = await monthly_document(facts=facts)
    review.approve(document, MONTHLY_SECTION, by="sam")

    result = render_pdf(document, tmp_path / "client.pdf", now=NOW)
    text = pdf_text(result.path)
    assert "No servers at this site." in text
    assert "Server lifecycle details" in text and result.omitted == []


async def test_the_internal_render_lists_devices_whose_os_and_node_class_disagree(
    tmp_path: Path,
) -> None:
    facts = monthly_facts()
    facts.tables[ROLE_DISAGREEMENTS] = table(
        ROLE_DISAGREEMENTS,
        ["Device", "Operating system", "NinjaOne node class", "Counted as"],
        [["ACME-WS02", "Windows 10 IoT Enterprise LTSC", "WINDOWS_SERVER", "workstation"]],
    )
    facts.tables[ROLE_DISAGREEMENTS].internal_only = True
    document = await monthly_document(facts=facts)
    review.approve(document, MONTHLY_SECTION, by="sam")

    internal = pdf_text(render_pdf(document, tmp_path / "i.pdf", mode="internal", now=NOW).path)
    assert "WINDOWS_SERVER" in internal and "Windows 10 IoT Enterprise LTSC" in internal
    client = pdf_text(render_pdf(document, tmp_path / "c.pdf", now=NOW).path)
    assert "WINDOWS_SERVER" not in client  # internal only
    # The deck's internal notes carry the same list under the slide with the counts.
    deck = render(document, tmp_path / "internal.pptx", mode="internal")
    assert "counted as a workstation" in deck_words(deck)
    assert "WINDOWS_SERVER" not in deck_words(render(document, tmp_path / "client.pptx"))


# --- people's names: the franchisee default on the PDF --------------------------------------


async def test_the_franchisee_default_names_people_on_the_pdf_and_demo_ignores_it(
    tmp_path: Path,
) -> None:
    document = await monthly_document(config=make_config(show_people_names=True))
    review.approve(document, MONTHLY_SECTION, by="sam")

    for mode in ("client", "internal"):
        text = pdf_text(render_pdf(document, tmp_path / f"{mode}.pdf", mode=mode, now=NOW).path)  # type: ignore[arg-type]
        # The contact chart, the lifecycle users and the person-named location, all named.
        assert "Dana Whitfield" in text and LOGINS[0] in text
        assert " ".join(LOCATIONS[1].split()) in " ".join(text.split())
        assert "Contact A" not in text and "User A" not in text and "withheld" not in text

    demo = pdf_text(render_pdf(document, tmp_path / "demo.pdf", mode="demo", now=NOW).path)
    assert substring_survivors(demo, real_name_terms(document)) == []

    # The review's override hides them again, and the private gate is armed with it.
    review.set_contact_names(document, False, by="sam")
    hidden = pdf_text(render_pdf(document, tmp_path / "hidden.pdf", now=NOW).path)
    assert "Dana Whitfield" not in hidden and "Contact A" in hidden and "User A" in hidden
