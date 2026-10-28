"""Who a deck is for: client, internal and demo renders; the Users line; the report allowlist."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pydantic import ValidationError

from app.reports.tbr import gate, review
from app.reports.tbr.audit import NullAudit
from app.reports.tbr.demo import CLIENT_NAME, SITE_NAME, STAMP, build_pseudonyms, pseudonymise
from app.reports.tbr.model import Item
from app.reports.tbr.narrative import report_search_tool, research_trends, source_allowed
from app.reports.tbr.render import TEMPLATE, default_output, render
from app.settings import DEFAULT_REPORT_DOMAINS, Settings
from tests.reports.conftest import NOW, ScriptedModel, make_document, make_facts
from tests.reports.test_tbr_rules import searched


def slide_for(result: Any, number: int) -> Any:
    """The rendered slide that was template slide `number`."""
    deck = Presentation(str(result.path))
    return list(deck.slides)[result.kept.index(number)]


def words(slide: Any) -> str:
    return " \n ".join([*gate.all_text(slide), *gate.chart_text(slide)])


# --- 1. render modes ----------------------------------------------------------------------------


def test_the_template_ships_speaker_notes_and_the_gate_refuses_them() -> None:
    slides = list(Presentation(str(TEMPLATE)).slides)
    noted = [n for n, slide in enumerate(slides, start=1) if gate.notes_text(slide)]
    assert noted == [1, 3, 6, 7]  # presenter instructions and AI prompts, not for a customer

    report = gate.check([(1, slides[0])], allowed_notes={})

    assert any("speaker notes" in problem for problem in report.problems)


async def test_a_client_deck_has_empty_notes_by_default(tmp_path: Path) -> None:
    document = await make_document()

    result = render(document, tmp_path / "deck.pptx")  # client is the default

    assert result.mode == "client"
    deck = Presentation(str(result.path))
    assert [gate.notes_text(slide) for slide in deck.slides] == [""] * len(result.kept)
    assert 1 in result.kept and 6 in result.kept  # slides whose template notes were cleared


async def test_a_client_deck_carries_the_engineers_talk_track_and_nothing_else(
    tmp_path: Path,
) -> None:
    document = await make_document()
    review.set_talk_track(document, "cybersecurity", " Open with the  patching result. ", by="sam")

    result = render(document, tmp_path / "deck.pptx", mode="client")

    assert gate.notes_text(slide_for(result, 14)) == "Open with the patching result."
    section = document.section("cybersecurity")
    assert (section.talk_track_by, section.talk_track) == ("sam", "Open with the patching result.")
    others = [n for n in result.kept if n != 14]
    assert all(gate.notes_text(slide_for(result, n)) == "" for n in others)
    assert "ninjaone." not in gate.notes_text(slide_for(result, 14))  # no fact ids


async def test_an_internal_deck_carries_the_fact_id_trace(tmp_path: Path) -> None:
    document = await make_document()
    review.set_talk_track(document, "cybersecurity", "Open with the patching result.", by="sam")

    result = render(document, tmp_path / "deck.pptx", mode="internal")

    notes = gate.notes_text(slide_for(result, 14))
    assert notes.startswith("Open with the patching result.")
    assert "Sources for this slide" in notes
    assert "ninjaone.patches.critical_pending" in notes and "Test definition of" in notes
    assert gate.notes_text(slide_for(result, 1)) == ""  # the title slide has nothing to trace
    assert not any("<<" in gate.notes_text(slide_for(result, n)) for n in result.kept)


async def test_the_trace_stays_on_the_review_whatever_is_rendered(tmp_path: Path) -> None:
    document = await make_document()
    before = document.model_dump(mode="json")

    for mode in ("client", "internal", "demo"):
        render(document, tmp_path / f"{mode}.pptx", mode=mode)  # type: ignore[arg-type]

    assert document.model_dump(mode="json") == before  # rendering never edits the review
    cited = [item for item in document.section("cybersecurity").all_items() if item.cites]
    known = set(document.facts.facts) | {finding.id for finding in document.findings}
    assert cited and all(set(item.cites) <= known for item in cited)  # every cite resolves


async def test_notes_that_are_not_the_talk_track_stop_a_client_render(tmp_path: Path) -> None:
    document = await make_document()
    result = render(document, tmp_path / "deck.pptx")
    deck = Presentation(str(result.path))
    slides = list(zip(result.kept, deck.slides, strict=True))
    slides[4][1].notes_slide.notes_text_frame.text = "Sources: ninjaone.devices.total"

    report = gate.check(slides, allowed_notes={})

    assert [p for p in report.problems if "not the engineer's talk track" in p]
    assert gate.check(slides, allowed_notes=None).passed  # an internal deck may hold a trace


async def test_speaker_notes_may_not_hold_template_placeholder_text() -> None:
    document = await make_document()

    with pytest.raises(review.ReviewError, match="placeholder"):
        review.set_talk_track(document, "cybersecurity", "<<insert industry here>>", by="sam")
    with pytest.raises(review.ReviewError, match="no section"):
        review.set_talk_track(document, "no-such-section", "Hello", by="sam")
    assert document.section("cybersecurity").talk_track is None


async def test_empty_text_clears_the_talk_track() -> None:
    document = await make_document()
    review.set_talk_track(document, "projects", "Mention the office move.", by="sam")

    review.set_talk_track(document, "projects", "  ", by="sam")

    assert document.section("projects").talk_track is None
    assert document.section("projects").talk_track_by is None


def test_the_file_name_says_which_deck_it_is(tmp_path: Path) -> None:
    import asyncio

    document = asyncio.run(make_document())

    client = default_output(document, tmp_path, now=NOW)
    internal = default_output(document, tmp_path, now=NOW, mode="internal")
    demo = default_output(document, tmp_path, now=NOW, mode="demo")

    assert client.name.startswith("acme-legal-tbr-2026-08-") and "internal" not in client.name
    assert internal.name.endswith("-internal.pptx")
    assert demo.name.startswith("sample-client-tbr-") and "acme" not in demo.name


# --- 2. no Users line ---------------------------------------------------------------------------


async def test_slide_five_has_no_users_line_and_no_contact_count(tmp_path: Path) -> None:
    document = await make_document()

    assert document.facts.get("site.users") is None
    assert document.facts.get("autotask.contacts.active") is None
    management = document.section("management")
    assert not [item for item in management.all_items() if "user" in (item.lead or "").lower()]
    assert "Microsoft 365" in management.reason

    shown = words(slide_for(render(document, tmp_path / "deck.pptx"), 5))
    assert "Users" not in shown and "Contacts" not in shown
    assert "44 Workstations" in shown  # the device counts are still there


async def test_a_review_stored_with_a_users_line_renders_without_it(tmp_path: Path) -> None:
    document = await make_document()
    document.section("management").items.insert(
        0, Item(role="bullet", lead="Users:", text="41", cites=["site.users"], author="code")
    )
    document.section("management").items.insert(
        0,
        Item(
            role="bullet",
            lead="Contacts:",
            text="41",
            cites=["autotask.contacts.active"],
            author="code",
        ),
    )

    shown = words(slide_for(render(document, tmp_path / "deck.pptx"), 5))

    assert "Users:" not in shown and "Contacts:" not in shown and "41" not in shown


def test_nothing_derives_a_finding_from_a_user_count() -> None:
    from tests.reports.test_tbr_rules import findings_for

    found = findings_for(make_facts(site__users=10, autotask__contacts__active=10))

    assert "inventory.more_machines_than_people" not in found
    assert not [finding_id for finding_id in found if "user" in finding_id]


# --- 3. the report allowlist --------------------------------------------------------------------


def test_reports_default_to_regulators_and_associations() -> None:
    domains = Settings().report_web_search_domains

    assert {"cisa.gov", "nist.gov", "ic3.gov", "ftc.gov", "hhs.gov", "sec.gov"} <= set(domains)
    assert {"americanbar.org", "aicpa-cima.com", "ada.org"} <= set(domains)
    assert DEFAULT_REPORT_DOMAINS.split(",")[0] == "cisa.gov"


@pytest.mark.parametrize(
    "value", ["cisa.gov,microsoft.com", "learn.microsoft.com", "https://www.ninjaone.com/blog", ""]
)
def test_a_vendor_domain_or_an_empty_list_is_refused(value: str) -> None:
    with pytest.raises(ValidationError):
        Settings(report_web_search_allowed_domains=value)


def test_the_report_list_is_separate_from_the_chat_list() -> None:
    settings = Settings(
        web_search_allowed_domains="learn.microsoft.com,ninjaone.com",
        web_search_max_uses=2,
        report_web_search_allowed_domains="cisa.gov, americanbar.org",
        report_web_search_max_uses=4,
    )

    tool = report_search_tool(settings)

    assert tool is not None
    assert tool["allowed_domains"] == ["cisa.gov", "americanbar.org"] and tool["max_uses"] == 4
    assert settings.web_search_domains == ["learn.microsoft.com", "ninjaone.com"]  # unchanged
    assert report_search_tool(Settings(web_search_enabled=False)) is None


def test_a_source_is_checked_by_host_not_by_substring() -> None:
    settings = Settings()

    assert source_allowed("https://www.cisa.gov/news/a", settings)
    assert source_allowed("https://csrc.nist.gov/pubs", settings)  # a subdomain of a listed one
    assert not source_allowed("https://www.microsoft.com/security/blog", settings)
    assert not source_allowed("https://cisa.gov.example.com/a", settings)
    assert not source_allowed("https://example.org/cisa.gov", settings)
    assert not source_allowed("not a url", settings)


async def test_a_trend_from_a_vendor_or_an_unlisted_site_is_dropped() -> None:
    urls = [
        "https://www.cisa.gov/a",
        "https://www.sonicwall.com/threat-report",
        "https://some-blog.example.com/post",
    ]
    model = ScriptedModel(
        searched(
            "law firm ransomware guidance",
            urls,
            [
                {"id": str(n), "text": f"Trend number {'abc'[n]}.", "cites": [], "url": url}
                for n, url in enumerate(urls)
            ],
        )
    )

    draft = await research_trends(
        "Legal Services",
        never_in_a_query=["Acme Legal"],
        model_client=model,  # type: ignore[arg-type]
        settings=Settings(web_search_enabled=True),
        audit=NullAudit(),
    )

    assert [item.url for item in draft.items] == ["https://www.cisa.gov/a"]
    assert sum("not a domain reports may cite" in p for p in draft.problems) == 2
    sent = model.requests[0]["tools"]
    search = next(t for t in sent if t.get("name") == "web_search")
    assert "cisa.gov" in search["allowed_domains"] and "sonicwall.com" not in search


# --- 4. demo decks ------------------------------------------------------------------------------

REAL = ["Acme Legal", "acme-legal", "Example IT Partners", "ACME-", "A. Person", "RECEPTION"]


async def demo_document() -> Any:
    document = await make_document()
    review.add_item(
        document, "accomplishments", "Rebuilt ACME-DC01 for Acme Legal after hours", by="sam"
    )
    review.set_talk_track(document, "projects", "Acme Legal asked about ACME-WS01.", by="sam")
    return document


async def test_a_demo_deck_names_no_one_real(tmp_path: Path) -> None:
    document = await demo_document()

    result = render(document, tmp_path / "demo.pptx", mode="demo")

    assert result.mode == "demo"
    deck = Presentation(str(result.path))
    everything = "\n".join(words(s) + "\n" + gate.notes_text(s) for s in deck.slides)
    assert not [name for name in REAL if name.casefold() in everything.casefold()]
    assert SITE_NAME in everything and CLIENT_NAME in everything
    assert "Rebuilt DEVICE-" in everything and "Contact A" in everything
    assert f"{SITE_NAME} asked about DEVICE-" in gate.notes_text(slide_for(result, 11))


async def test_the_same_name_gets_the_same_pseudonym_on_every_slide(tmp_path: Path) -> None:
    document = await demo_document()
    names = build_pseudonyms(document)
    client = render(document, tmp_path / "client.pptx")
    demo = render(document, tmp_path / "demo.pptx", mode="demo")
    assert client.kept == demo.kept

    # Wherever a real name stood, its one pseudonym stands: same slide, same number of times.
    seen = 0
    for number in client.kept:
        real_text, demo_text = words(slide_for(client, number)), words(slide_for(demo, number))
        for real, fake in names.exact.items():
            if not real.startswith("ACME-"):
                continue
            pattern = rf"(?<![A-Za-z0-9]){re.escape(real)}(?![A-Za-z0-9])"
            count = len(re.findall(pattern, real_text))
            seen += count
            assert len(re.findall(rf"{re.escape(fake)}(?!\d)", demo_text)) == count, (number, real)
    assert seen >= 4  # the names really are on more than one slide
    assert len(set(names.exact.values())) == len(names.exact)  # and no two share a pseudonym
    assert build_pseudonyms(document).exact == names.exact  # the mapping is deterministic


def _figures(text: str, drop: list[str]) -> list[str]:
    for name in sorted(drop, key=len, reverse=True):
        text = text.replace(name, " ")
    return re.findall(r"\d+(?:[.,]\d+)*%?", text)


def _chart_values(slide: Any) -> list[list[float]]:
    found = []
    for shape in slide.shapes:
        if getattr(shape, "has_chart", False) and shape.has_chart:
            for plot in shape.chart.plots:
                found.extend([list(series.values) for series in plot.series])
    return found


async def test_a_demo_deck_keeps_every_number(tmp_path: Path) -> None:
    document = await demo_document()
    names = build_pseudonyms(document)
    client = render(document, tmp_path / "client.pptx")
    demo = render(document, tmp_path / "demo.pptx", mode="demo")

    compared = 0
    for number in client.kept:
        real_slide, demo_slide = slide_for(client, number), slide_for(demo, number)
        real = _figures(words(real_slide), list(names.exact))
        fake = _figures(words(demo_slide), list(names.exact.values()))
        assert real == fake, number
        assert _chart_values(real_slide) == _chart_values(demo_slide), number
        compared += len(real)
    assert compared > 30


async def test_every_demo_slide_is_stamped_and_no_other_deck_is(tmp_path: Path) -> None:
    document = await make_document()

    demo = Presentation(str(render(document, tmp_path / "demo.pptx", mode="demo").path))
    client = Presentation(str(render(document, tmp_path / "client.pptx").path))

    for slide in demo.slides:
        stamps = [s for s in slide.shapes if s.name == "Sample data stamp"]
        assert len(stamps) == 1 and stamps[0].text_frame.text == STAMP
        assert stamps[0].left > demo.slide_width * 0.8 and stamps[0].top < demo.slide_height * 0.1
    assert not any(STAMP in words(slide) for slide in client.slides)


async def test_the_gate_stops_a_demo_deck_with_a_real_name_left_in_it(tmp_path: Path) -> None:
    document = await demo_document()
    review.set_contact_names(document, True, by="sam")  # a client deck with every real name
    names = build_pseudonyms(document)
    client = render(document, tmp_path / "client.pptx")
    deck = Presentation(str(client.path))
    slides = list(zip(client.kept, deck.slides, strict=True))

    report = gate.check(slides, real_names=names.survivors)

    survived = " ".join(report.problems)
    assert "a real name survived in a demo deck" in survived
    assert "'Acme Legal'" in survived and "'ACME-DC01'" in survived
    assert "'A. Person'" in survived  # found in a chart's category labels, not in slide text


async def test_a_device_name_that_is_also_a_word_is_matched_as_a_name_only() -> None:
    document = await make_document()
    review.add_item(document, "accomplishments", "Moved RECEPTION to the reception desk", by="sam")

    copy, names = pseudonymise(document)

    text = copy.section("accomplishments").engineer_items[-1].text
    assert text == f"Moved {names.exact['RECEPTION']} to the reception desk"
    assert copy.facts.names == {}  # the copy does not carry the list of real names
    assert document.facts.names["devices"]  # and the stored review still does
    assert copy.generated_by.startswith("Engineer ")


# --- what the gate refuses in speaker notes -----------------------------------------------------

ORDINARY = (
    "My talk track: insert the new drive first, then list recent changes. "
    "Examples provided on request; list specific servers only if asked."
)


async def test_an_engineers_notes_may_use_the_phrases_a_slide_may_not(tmp_path: Path) -> None:
    document = await make_document()
    assert not gate.note_problems(ORDINARY)

    review.set_talk_track(document, "cybersecurity", ORDINARY, by="sam")
    result = render(document, tmp_path / "deck.pptx")  # the gate reads these notes

    assert result.gate.passed
    assert gate.notes_text(slide_for(result, 14)) == ORDINARY
    # The same words on a slide are still an instruction someone left behind.
    with pytest.raises(review.ReviewError, match="placeholder"):
        review.add_item(document, "accomplishments", "List recent wins", by="sam")


def test_every_note_the_template_ships_is_refused() -> None:
    slides = list(Presentation(str(TEMPLATE)).slides)
    for number in (1, 3, 6, 7):
        report = gate.check([(number, slides[number - 1])])  # any mode: no allowed_notes
        found = [p for p in report.problems if "the template's own notes text" in p]
        assert found, number


def test_one_line_of_the_templates_notes_inside_other_text_is_refused() -> None:
    line = "Identify 3 industry-specific technology trends leadership must address"
    assert (
        line.casefold() in gate.notes_text(list(Presentation(str(TEMPLATE)).slides)[2]).casefold()
    )
    fragments = gate.template_note_fragments()

    pasted = "Remember to mention the office move.\nFormatting:\n•  Slide-ready with  no rewriting."
    problems = gate.note_problems(pasted)

    assert "slide ready with no rewriting" in fragments
    assert len(problems) == 1 and "slide ready with no rewriting" in problems[0]
    assert all(len(fragment.split()) == gate.MIN_FRAGMENT_WORDS for fragment in fragments)
    # Half a sentence, mid-line, inside the engineer's own words.
    assert gate.note_problems(f"First, {line[12:]}, then the budget.")
    assert not gate.note_problems("Rules: no technical jargon. Guidelines: keep it short.")


@pytest.mark.parametrize(
    "text",
    ["Replace XX laptops", "Ask ABC Company", "Due in 202x", "{insert your opinion}", "<<client>>"],
)
async def test_a_fill_in_mark_in_the_notes_is_refused(text: str) -> None:
    document = await make_document()

    assert gate.note_problems(text)
    with pytest.raises(review.ReviewError, match="placeholder mark"):
        review.set_talk_track(document, "projects", text, by="sam")
    assert not gate.note_problems("The fix went to the xx-large VM pool")  # XX is as written
