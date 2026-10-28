"""Every model-written slide, rendered at its budget's maximum and at one item.

The budget in `layout.SLIDE_BUDGETS` is a promise that a slide the model
writes to it fits its placeholder. This is where the promise is checked: each
budgeted section is filled with prose at exactly its limits, with the longest
lead the code can put on it, rendered through the real renderer into the real
template, and then the saved file is read back and measured against the
placeholder's own height. At one item the same slides must take the template's
size, so a slide with little on it is not set at the size the fullest one
needs.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Emu

from app.reports.tbr import layout
from app.reports.tbr.layout import HELD_LINE, SLIDE_BUDGETS, SlideBudget, fit, shorten
from app.reports.tbr.model import Item, Section
from app.reports.tbr.narrative import (
    MAX_SLOT_CHARS,
    Brief,
    Slot,
    _brief_json,
    printed_length,
    problems_with,
)
from app.reports.tbr.render import items_as_paragraphs, render
from app.reports.tbr.sections import RECOMMENDATION_CHARS
from tests.reports.conftest import make_document, make_facts

PROSE = (
    "Seven devices have not checked in for fourteen days or more, and the oldest has been "
    "silent since the second of September, which leaves their patch state unknown and their "
    "antivirus definitions older than the policy allows for any managed workstation. "
) * 8
#: The longest headline `findings.py` writes, and the longest service name in the
#: default catalogue, as the leads the renderer puts in bold before the model's text.
LONGEST_HEADLINE = "Password and lockout requests are a steady share of the helpdesk load:"
LONGEST_SERVICE = (
    "Secure Access Service Edge + Network Data Loss Prevention "
    + chr(0x2013)  # the catalogue's en dash
    + " Secure Login and Permission Management:"
)
MODEL_WRITTEN = {"management", "dashboard", "takeaways", "projects", "cybersecurity", "services"}


def prose(chars: int) -> str:
    return PROSE[:chars].rsplit(" ", 1)[0] + "."


def model(text: str, slot: str, **fields: object) -> Item:
    return Item(text=text, cites=["x"], author="model", slot=slot, **fields)  # type: ignore[arg-type]


def code(text: str, slot: str | None = None, **fields: object) -> Item:
    return Item(text=text, cites=["x"], author="code", slot=slot, **fields)  # type: ignore[arg-type]


def items_at(budget: SlideBudget, count: int) -> list[Item]:
    """A section's items with `count` model-written items at the character budget,
    plus every code-written line the section puts beside them."""
    chars, key = budget.chars_per_item, budget.section
    items: list[Item] = []
    if key == "cybersecurity":
        for i in range(count):
            items.append(
                model(prose(chars), f"f{i}/observation", role="observation", lead=LONGEST_HEADLINE)
            )
            items.append(
                model(
                    prose(min(chars, RECOMMENDATION_CHARS)),
                    f"f{i}/recommendation",
                    role="recommendation",
                    lead="Recommendation:",
                )
            )
        items.append(
            code(
                "A further 1,234 critical findings are suspected but not yet verified by a "
                "scan, and are not counted above.",
                role="paragraph",
            )
        )
    elif key == "services":
        items += [
            model(prose(chars), f"service/s{i}", role="bullet", lead=LONGEST_SERVICE)
            for i in range(count)
        ]
    elif key == "projects":
        for i in range(count):
            items += [
                code(prose(60), f"project-{i}/name", role="headline"),
                code(
                    "Acme Legal is in the Implementation phase of the project, 45% complete.",
                    f"project-{i}/phase",
                    role="paragraph",
                ),
                model(prose(chars), f"project-{i}/scope", role="paragraph"),
                code(
                    "The project was planned to complete by the end of August 2026 and is past "
                    "that date; a revised date needs to be agreed.",
                    f"project-{i}/due",
                    role="paragraph",
                ),
            ]
    elif key == "takeaways":
        items += [
            model(
                prose(chars),
                f"f{i}/row",
                role="paragraph",
                lead="A headline of at most ten words goes here",
            )
            for i in range(count)
        ]
    elif key == "trends":
        items += [
            model(prose(chars), f"trend/{i}", role="bullet", url="https://www.cisa.gov/x")
            for i in range(count)
        ]
    elif key == "management":
        for lead, label in (
            ("123", "Workstations and Laptops"),
            ("4", "Physical Servers"),
            ("12", "Virtual Servers"),
            ("3", "Locations"),
        ):
            items.append(code(label, lead=lead, role="bullet"))
        items.append(
            code(
                "Based on a technical review of the last three months, all segments of Example "
                "IT Partners of a Long Placename's services offering have been reviewed and "
                "found to need attention.",
                role="paragraph",
            )
        )
        items += [model(prose(chars), "improvements", role="paragraph") for _ in range(count)]
    elif key == "dashboard":
        items += [model(prose(chars), f"f{i}/observation", role="bullet") for i in range(count)]
    return items


async def rendered_with(tmp_path: Path, count_of: dict[str, int]) -> tuple[Presentation, Any]:
    """A deck with every budgeted section filled at `count_of[section]` items."""
    document = await make_document(facts=make_facts())
    by_key = {section.key: section for section in document.sections}
    for key, count in count_of.items():
        budget = SLIDE_BUDGETS[key]
        section = by_key[key]
        section.status = "filled"
        section.narrative = "model"
        section.items = items_at(budget, count)
        section.engineer_items = []
        section.requires_approval = False
    result = render(document, tmp_path / "deck.pptx", mode="internal")
    assert result.gate.passed
    deck = Presentation(str(result.path))
    return deck, result


def body_of(slide: object) -> object:
    return next(s for s in slide.placeholders if s.placeholder_format.idx == 1)  # type: ignore[attr-defined]


def slide_by_title(deck: Presentation, number: int, kept: list[int]) -> object:
    return deck.slides[kept.index(number)]


def assert_fits(shape: object, *, what: str) -> float:
    box = layout.box_of(shape)
    measured = layout.measure_shape(shape, box)
    assert measured.height <= box.height, (
        f"{what}: {measured.height:.0f}pt of text in a {box.height:.0f}pt placeholder"
    )
    return measured.height / box.height


@pytest.mark.parametrize("key", sorted(SLIDE_BUDGETS))
async def test_a_slide_at_its_budget_fits_its_placeholder(tmp_path: Path, key: str) -> None:
    budget = SLIDE_BUDGETS[key]
    deck, result = await rendered_with(tmp_path, {key: budget.max_items})
    slide = slide_by_title(deck, budget.slide, result.kept)
    if key == "takeaways":
        frame = next(s for s in slide.shapes if getattr(s, "has_table", False) and s.has_table)
        rows = list(frame.table.rows)
        assert len(rows) == budget.max_items
        for index, row in enumerate(rows):
            cell = layout.TextBox(
                width=Emu(frame.table.columns[1].width).inches * layout.POINTS_PER_INCH
                - 2 * layout.DEFAULT_INSET_X,
                height=Emu(row.height).inches * layout.POINTS_PER_INCH - 2 * layout.DEFAULT_INSET_Y,
                indents=(0.0,),
            )
            measured = layout.measure(layout.read_paragraphs(row.cells[1]), cell)
            assert measured.height <= cell.height, f"row {index + 1} overflows its cell"
        # And the table itself ends above the footer strip.
        bottom = Emu(frame.top + sum(r.height for r in rows)).inches
        assert bottom <= 5.02
        return
    body = body_of(slide)
    assert_fits(body, what=f"slide {budget.slide} at {budget.max_items} items")
    # Nothing was dropped to make it fit: the budget is the promise, not the fitter.
    text = " ".join(p.text for p in layout.read_paragraphs(body))
    assert HELD_LINE.split("}")[-1].strip() not in text
    # Every budgeted item is on the slide, whole.
    expected = items_at(budget, budget.max_items)
    for item in expected:
        if item.author == "model":
            assert item.text in text, f"{item.slot} was shortened or dropped"


@pytest.mark.parametrize("key", sorted(SLIDE_BUDGETS))
async def test_a_slide_with_one_item_takes_the_largest_size(tmp_path: Path, key: str) -> None:
    budget = SLIDE_BUDGETS[key]
    deck, result = await rendered_with(tmp_path, {key: 1})
    slide = slide_by_title(deck, budget.slide, result.kept)
    if key == "takeaways":
        frame = next(s for s in slide.shapes if getattr(s, "has_table", False) and s.has_table)
        rows = list(frame.table.rows)
        assert len(rows) == 1
        sizes = {p.size for p in layout.read_paragraphs(rows[0].cells[1])}
        assert max(sizes) >= budget.sizes[0]
        return
    body = body_of(slide)
    paragraphs = layout.read_paragraphs(body)
    if key == "dashboard":
        paragraphs = paragraphs[1:]  # the template's own "Observations" heading
    model_sizes = {p.size for p in paragraphs if p.size and p.level == 0}
    assert budget.sizes[0] in model_sizes, f"one item was set at {model_sizes}, not the largest"
    assert_fits(body, what=f"slide {budget.slide} at one item")


def test_the_limit_counts_the_text_as_it_prints_not_the_token_names() -> None:
    facts = make_facts()
    token = "{{ninjaone.patches.failed}}"  # prints as "14"
    text = f"{token} patch installations failed, which is {token} too many."
    assert printed_length(text, facts) == len(text.replace(token, "14"))
    # Under the limit once printed, over it as written: accepted.
    limit = printed_length(text, facts)
    assert len(text) > limit
    assert not [
        p
        for p in problems_with(
            text,
            ["ninjaone.patches.failed"],
            allowed={"ninjaone.patches.failed"},
            facts=facts,
            limit=limit,
        )
        if "longer than" in p
    ]
    over = problems_with(
        text,
        ["ninjaone.patches.failed"],
        allowed={"ninjaone.patches.failed"},
        facts=facts,
        limit=limit - 1,
    )
    assert any("once its tokens print" in p and "cut at least 1 " in p for p in over)


def test_every_model_written_section_has_a_budget() -> None:
    assert MODEL_WRITTEN | {"trends"} == set(SLIDE_BUDGETS)


def test_the_brief_tells_the_model_its_budget_up_front() -> None:
    budget = SLIDE_BUDGETS["services"]
    brief = Brief(
        key="services",
        title="Recommended Services",
        purpose="p",
        fact_ids=[],
        findings=[],
        slots=[
            Slot(id="service/a", guidance="g"),
            Slot(id="service/b", guidance="g", max_chars=40),
        ],
        budget=budget,
    )
    text = _brief_json(brief, make_facts())
    assert f'"max_chars": {budget.chars_per_item}' in text
    assert '"max_chars": 40' in text
    assert f'"slide_max_chars": {budget.chars_per_slide}' in text
    assert brief.limit_for(brief.slots[0]) == budget.chars_per_item
    assert brief.limit_for(brief.slots[1]) == 40
    unbudgeted = Brief(key="x", title="x", purpose="p", fact_ids=[], findings=[], slots=[])
    assert unbudgeted.limit_for(Slot(id="s", guidance="g")) == MAX_SLOT_CHARS


# --- the fitter on its own --------------------------------------------------------------------


def _box(height: float = 120.0) -> layout.TextBox:
    return layout.TextBox(width=400.0, height=height)


def _budget(**fields: object) -> SlideBudget:
    base = dict(
        section="t",
        slide=99,
        max_items=4,
        chars_per_item=80,
        chars_per_slide=320,
        what="thing",
        sizes=(16, 14, 12),
    )
    return SlideBudget(**{**base, **fields})  # type: ignore[arg-type]


def test_fit_takes_the_largest_size_that_fits_before_touching_the_text() -> None:
    items = [model(prose(80), f"a{i}") for i in range(2)]
    result = fit(items, _box(120), budget=_budget(), paragraphs_of=items_as_paragraphs)
    assert not result.changed
    assert result.size == max(
        size for size in (16, 14, 12) if layout.fits(items_as_paragraphs(items, size), _box(120))
    )


def test_fit_drops_whole_groups_from_the_end_and_says_so() -> None:
    items = []
    for i in range(4):
        items.append(model(prose(120), f"f{i}/observation", role="observation", lead="Finding:"))
        items.append(model(prose(60), f"f{i}/recommendation", role="recommendation"))
    result = fit(items, _box(110), budget=_budget(), paragraphs_of=items_as_paragraphs)
    assert result.dropped
    assert result.size == 12  # only at the floor is anything dropped
    # Pairs go together: never an observation without its recommendation.
    slots = [item.slot for item in result.items if item.slot]
    kept = {slot.split("/")[0] for slot in slots}
    for finding in kept:
        assert f"{finding}/observation" in slots and f"{finding}/recommendation" in slots
    # The most serious (first) survive and the line at the end says how many did not.
    assert slots[0] == "f0/observation"
    assert result.items[-1].author == "code"
    assert result.items[-1].text.startswith(str(len(result.dropped) // 2))
    assert layout.fits(items_as_paragraphs(result.items, result.size), _box(110))


def test_fit_never_drops_or_shortens_a_pinned_item() -> None:
    caveat = code("Suspected findings are not counted above.", role="paragraph")
    items = [model(prose(200), f"t{i}") for i in range(3)] + [caveat]
    result = fit(items, _box(60), budget=_budget(), paragraphs_of=items_as_paragraphs)
    assert caveat in result.items
    assert result.items.index(caveat) == len(result.items) - 1  # after the held line
    assert all(item.slot for item in result.dropped)


def test_fit_shortens_at_sentence_boundaries_only_after_dropping_to_the_floor() -> None:
    long = "First sentence here. Second sentence follows it. Third and last one."
    items = [model(long, "only")]
    box = layout.TextBox(width=300.0, height=28.0)  # one line at the floor, and no more
    result = fit(items, box, budget=_budget(chars_per_item=50), paragraphs_of=items_as_paragraphs)
    assert not result.dropped  # one group is the floor
    assert result.shortened == 1
    assert result.items[0].text == "First sentence here. Second sentence follows it."
    assert items[0].text == long  # the stored item is untouched


def test_shorten_keeps_a_first_sentence_that_is_itself_over_the_limit() -> None:
    assert shorten("One very long single sentence with no break.", 10) == (
        "One very long single sentence with no break."
    )
    assert shorten("Short. Also short.", 100) == "Short. Also short."


def test_a_section_of_the_wrong_shape_is_not_fitted() -> None:
    section = Section(key="accomplishments", slide=8, title="A", status="filled", reason="")
    assert layout.budget_for(section.key) is None
