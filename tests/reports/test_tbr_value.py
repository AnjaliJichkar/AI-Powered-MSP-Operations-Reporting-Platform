"""Value delivered on the Accomplishments slide, and the business hours behind one figure."""

from __future__ import annotations

import datetime as dt
import uuid
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from pptx import Presentation
from pydantic import ValidationError

from app.reports.tbr import gate, review
from app.reports.tbr.audit import NullAudit
from app.reports.tbr.collect import OUT_OF_HOURS, Collector
from app.reports.tbr.config import (
    BusinessHours,
    FranchiseeConfig,
    TimeZoneUnavailable,
    resolve_zone,
)
from app.reports.tbr.render import render
from app.reports.tbr.sections import VALUE_DELIVERED
from tests.reports.conftest import PERIOD, make_document, make_facts

UTC = dt.UTC


def slide_text(result: Any, number: int) -> list[str]:
    deck = Presentation(str(result.path))
    return gate.all_text(list(deck.slides)[result.kept.index(number)])


# --- the block ----------------------------------------------------------------------------------


async def test_only_the_figures_the_connectors_supplied_are_listed() -> None:
    document = await make_document()  # no alert history, no business hours: two of four

    section = document.section("accomplishments")
    block = [item for item in section.items if (item.slot or "").startswith("value/")]

    assert [(item.role, item.lead, item.text) for item in block] == [
        ("headline", None, "Value delivered in August 2026"),
        ("bullet", "310", "Windows patches installed"),
        ("bullet", "115", "tickets resolved"),
    ]
    assert all(item.author == "code" for item in block)  # counted and worded in code
    assert block[1].cites == ["ninjaone.patches.installed"]
    assert "left out: monitoring alerts resolved, tickets resolved outside business" in (
        section.reason
    )


async def test_all_four_figures_each_traced_to_a_fact_with_a_definition() -> None:
    facts = make_facts(
        ninjaone__alerts__resolved=27,
        autotask__tickets__completed_out_of_hours=9,
        report__business_hours="Monday to Friday, 08:00 to 17:00",
    )
    document = await make_document(facts=facts)

    block = [i for i in document.section("accomplishments").items if i.role == "bullet"][:4]

    assert [f"{item.lead} {item.text}" for item in block] == [
        "310 Windows patches installed",
        "27 monitoring alerts resolved",
        "115 tickets resolved",
        "9 tickets resolved outside business hours (Monday to Friday, 08:00 to 17:00)",
    ]
    assert block[3].cites == ["autotask.tickets.completed_out_of_hours", "report.business_hours"]
    for fact_id, _, _ in VALUE_DELIVERED:
        assert document.facts.facts[fact_id].definition


async def test_a_zero_or_a_missing_figure_has_no_line_and_one_is_singular() -> None:
    facts = make_facts(ninjaone__alerts__resolved=0, autotask__tickets__completed=1)
    del facts.facts["ninjaone.patches.installed"]

    document = await make_document(facts=facts)

    lines = [i for i in document.section("accomplishments").items if i.role == "bullet"]
    assert [f"{i.lead} {i.text}" for i in lines if (i.slot or "").startswith("value/")] == [
        "1 ticket resolved"
    ]


async def test_with_nothing_to_count_the_slide_still_waits_for_the_engineer() -> None:
    facts = make_facts()
    for fact_id, _, _ in VALUE_DELIVERED:
        facts.facts.pop(fact_id, None)
    del facts.tables["autotask.projects.completed.list"]

    section = (await make_document(facts=facts)).section("accomplishments")

    assert (section.status, section.items) == ("awaiting_engineer", [])


async def test_the_block_comes_first_on_the_slide_and_the_rest_under_its_own_heading(
    tmp_path: Path,
) -> None:
    document = await make_document()
    review.add_item(document, "accomplishments", "Moved the finance share", by="sam")

    shown = slide_text(render(document, tmp_path / "deck.pptx"), 8)

    assert shown[1:] == [
        "Value delivered in August 2026",
        "310 Windows patches installed",
        "115 tickets resolved",
        "Also in this period",
        "Completed project: UPS swap (18 Aug 2026).",
        "Moved the finance share",
    ]


# --- business hours -----------------------------------------------------------------------------


def test_business_hours_are_local_start_inclusive_end_exclusive() -> None:
    hours, zone = BusinessHours(timezone="UTC-05:00"), resolve_zone("UTC-05:00")
    monday = dt.datetime(2026, 8, 3, tzinfo=UTC)

    assert not hours.contains(monday.replace(hour=12, minute=59), zone)  # 07:59 local
    assert hours.contains(monday.replace(hour=13), zone)  # 08:00 local
    assert hours.contains(monday.replace(hour=21, minute=59), zone)  # 16:59 local
    assert not hours.contains(monday.replace(hour=22), zone)  # 17:00 local
    # 02:00 UTC on Saturday is 21:00 on Friday locally: outside by the hour, not the day.
    assert not hours.contains(dt.datetime(2026, 8, 8, 2, tzinfo=UTC), zone)
    assert not hours.contains(dt.datetime(2026, 8, 8, 15, tzinfo=UTC), zone)  # Saturday 10:00
    assert hours.describe() == "Monday to Friday, 08:00 to 17:00"
    assert BusinessHours(days=[0, 2, 5]).describe().startswith("Monday, Wednesday, Saturday")


def test_a_zone_is_a_fixed_offset_or_a_name_and_a_bad_one_says_why() -> None:
    assert resolve_zone("UTC").utcoffset(None) == dt.timedelta(0)
    assert resolve_zone("utc+05:30").utcoffset(None) == dt.timedelta(hours=5, minutes=30)
    assert resolve_zone("UTC-6").utcoffset(None) == dt.timedelta(hours=-6)
    with pytest.raises(TimeZoneUnavailable, match="tzdata"):
        resolve_zone("Nowhere/Invented")
    with pytest.raises(ValidationError):
        BusinessHours(days=[7])


def collector(**hours: Any) -> Collector:
    return Collector(
        client_id=str(uuid.uuid4()),
        site_id=str(uuid.uuid4()),
        period=PERIOD,
        config=FranchiseeConfig(display_name="Example IT", business_hours=BusinessHours(**hours)),
        audit=NullAudit(),
    )


def completed(*moments: dt.datetime | None) -> list[Any]:
    return [SimpleNamespace(completed_date=moment) for moment in moments]


def test_tickets_completed_outside_hours_are_counted_in_the_customers_time() -> None:
    run = collector(timezone="UTC-05:00")
    tickets = completed(
        dt.datetime(2026, 8, 3, 15, tzinfo=UTC),  # Monday 10:00 local: inside
        dt.datetime(2026, 8, 4, 2, tzinfo=UTC),  # Monday 21:00 local: outside
        dt.datetime(2026, 8, 8, 15, tzinfo=UTC),  # Saturday: outside
        None,  # no completed time: cannot be placed, so not counted either way
    )

    run._out_of_hours(tickets)

    fact = run.facts.facts["autotask.tickets.completed_out_of_hours"]
    assert fact.value == 2 and fact.period_bound
    assert "Monday to Friday, 08:00 to 17:00" in fact.definition and "UTC-05:00" in fact.definition
    assert "when the ticket was completed" in fact.definition
    assert run.facts.facts["report.business_hours"].display == "Monday to Friday, 08:00 to 17:00"


@pytest.mark.parametrize(
    ("hours", "why"),
    [({}, "No business-hours time zone is set"), ({"timezone": "Nowhere/Invented"}, "tzdata")],
)
def test_without_a_usable_time_zone_the_figure_is_left_out(hours: dict[str, str], why: str) -> None:
    run = collector(**hours)

    run._out_of_hours(completed(dt.datetime(2026, 8, 8, 15, tzinfo=UTC)))

    assert "autotask.tickets.completed_out_of_hours" not in run.facts.facts  # never guessed
    assert [gap.what for gap in run.facts.gaps] == [OUT_OF_HOURS]
    assert why in run.facts.gaps[0].why
